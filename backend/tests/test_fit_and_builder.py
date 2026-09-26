"""Tests for the JD fit filter and the resume builder.

Run from the project root:

    pytest backend/tests -q

Covers the tech-domain gate, the fit verdict and its knockouts, the builder's
two consent modes, the honesty guarantee that no number is invented, and the
export renderers.
"""

from __future__ import annotations

import io
import re
import sys
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from services.ai_service import _introduces_metric  # noqa: E402
from services.analysis_service import get_analysis_service  # noqa: E402
from services.fit_evaluator import (  # noqa: E402
    COVERAGE_KNOCKOUT,
    VERDICT_GOOD,
    VERDICT_NONE,
    VERDICT_STRONG,
)
from services.jd_parser import (  # noqa: E402
    MIN_TECH_SKILLS,
    NonTechJobDescriptionError,
    get_jd_parser,
)
from services.knowledge_base import JobRequirement, get_knowledge_base  # noqa: E402
from services.resume_builder import (  # noqa: E402
    MODE_AUTO_ADD,
    MODE_CONFIRMED,
    BuildOptions,
    get_resume_builder,
)
from services.resume_export import render, strip_placeholders  # noqa: E402
from services.resume_parser import get_resume_parser  # noqa: E402

SAMPLES = BACKEND.parent / "samples"
_NUMBER_RE = re.compile(r"\d+(?:[.,]\d+)?%?")


# ---------------------------------------------------------------------------
# Job descriptions used across the tests
# ---------------------------------------------------------------------------

SALES_JD = """Senior Sales Representative - Enterprise Accounts
We are looking for a driven sales professional to join our growing sales team.
Responsibilities:
- Own the full sales cycle from prospecting to close for enterprise customers
- Build and maintain strong customer relationships and upsell existing accounts
- Hit and exceed quarterly sales quota and revenue targets
- Work closely with marketing to develop lead generation campaigns
Requirements:
- 5+ years of B2B sales or account management experience
- Excellent communication and negotiation skills
- Experience with CRM tools and a track record of exceeding quota
- Bachelor's degree in Business, Marketing or a related field
"""

NURSING_JD = """Registered Nurse - Medical Surgical Unit
Our hospital is seeking a compassionate registered nurse for patient care duties.
Responsibilities:
- Provide direct patient care and administer medication per physician orders
- Document patient records accurately and update the care team each shift
- Support families and coordinate with the nursing staff on shift handover
Requirements:
- Active RN license and a Bachelor of Science in Nursing
- 2+ years of nursing experience in an acute care setting
- Excellent communication skills and strong attention to detail
"""

ML_JD = """Machine Learning Engineer
We are hiring a machine learning engineer to build and deploy production models.
Responsibilities:
- Design, train and evaluate deep learning models using PyTorch and TensorFlow
- Build data pipelines with Spark and Airflow for large scale training data
- Deploy models to AWS SageMaker and monitor them in production with MLOps tooling
Basic Qualifications:
- 2+ years of experience building machine learning systems in Python
- Strong software engineering fundamentals, Docker and Kubernetes experience
- Experience with model deployment and feature engineering
Preferred Qualifications:
- Experience with Terraform and model monitoring
"""

SENIOR_ML_JD = """Senior Machine Learning Engineer
We need a senior ML engineer to own model development end to end.
Basic Qualifications:
- 6+ years of industry experience building machine learning systems
- Expert Python, PyTorch, Docker, Kubernetes and AWS SageMaker experience
- Experience with MLOps, model deployment and distributed training at scale
Responsibilities:
- Train and deploy deep learning models serving millions of requests
- Build data pipelines with Spark and Airflow
"""


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def sample_text() -> str:
    return (SAMPLES / "sample_resume.txt").read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def strong_text() -> str:
    return (SAMPLES / "strong_resume.txt").read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def parsed(sample_text):
    return get_resume_parser().parse_text(sample_text, "sample_resume.txt")


@pytest.fixture(scope="module")
def parsed_strong(strong_text):
    return get_resume_parser().parse_text(strong_text, "strong_resume.txt")


@pytest.fixture(scope="module")
def requirement():
    return get_knowledge_base().build_requirement(
        "amazon", "machine-learning-engineer", "entry"
    )


@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient

    import main

    with TestClient(main.app) as test_client:
        yield test_client


def _jd_analysis(text: str):
    """Parse a JD and return (requirement, analysis)."""
    return get_jd_parser().parse(text)


# ---------------------------------------------------------------------------
# The tech-domain gate
# ---------------------------------------------------------------------------


class TestTechDomainGate:
    @pytest.mark.parametrize("jd", [SALES_JD, NURSING_JD], ids=["sales", "nursing"])
    def test_non_tech_postings_are_refused(self, jd):
        parser = get_jd_parser()
        domain = parser.assess_domain(jd)
        assert domain["is_tech"] is False
        assert domain["tech_skill_count"] < MIN_TECH_SKILLS
        assert domain["non_tech_signals"], "the refusal must name what triggered it"

        with pytest.raises(NonTechJobDescriptionError) as excinfo:
            parser.parse(jd)
        assert "doesn't look like a tech job description" in str(excinfo.value)
        assert excinfo.value.domain["is_tech"] is False

    @pytest.mark.parametrize("jd", [ML_JD, SENIOR_ML_JD], ids=["ml", "senior-ml"])
    def test_tech_postings_pass(self, jd):
        domain = get_jd_parser().assess_domain(jd)
        assert domain["is_tech"] is True
        assert domain["tech_skill_count"] >= MIN_TECH_SKILLS
        requirement, analysis = get_jd_parser().parse(jd)
        assert requirement.by_tier("required")
        assert analysis["domain"]["is_tech"] is True

    def test_a_tech_posting_mentioning_customers_still_passes(self):
        """One non-tech word must not veto a posting full of technical skills."""
        jd = ML_JD + "\n- Partner with sales and marketing on customer support tooling\n"
        assert get_jd_parser().assess_domain(jd)["is_tech"] is True

    def test_data_analyst_counts_as_tech_only_with_a_technical_tool(self):
        with_tools = """Data Analyst
        Join our analytics team to turn raw data into decisions for the business.
        Responsibilities:
        - Write complex SQL queries against the data warehouse to answer questions
        - Build dashboards in Tableau and present findings to senior stakeholders
        Requirements:
        - 2+ years of analytics experience with advanced SQL and Python
        - Experience with Excel, statistics and A/B testing methodology
        """
        assert get_jd_parser().assess_domain(with_tools)["is_tech"] is True

    def test_degree_requirement_is_extracted(self):
        jd = ML_JD.replace(
            "Basic Qualifications:",
            "Basic Qualifications:\n- Master's degree in Computer Science or a related technical field",
        )
        _, analysis = _jd_analysis(jd)
        degree = analysis["degree_required"]
        assert degree["level"] == "master"
        assert degree["stated"] is True
        assert degree["field"] == "Computer Science"

    def test_ms_office_is_not_read_as_a_masters_degree(self):
        jd = ML_JD + "\n- Comfortable with MS Office and Google Workspace\n"
        _, analysis = _jd_analysis(jd)
        assert analysis["degree_required"]["level"] != "master"


class TestTechDomainGateAPI:
    def test_sales_jd_returns_422_with_code(self, client, sample_text):
        response = client.post(
            "/api/analyze/job-description",
            json={"resume_text": sample_text, "job_description": SALES_JD, "persist": False},
        )
        assert response.status_code == 422
        body = response.json()
        assert body["code"] == "non_tech_jd"
        assert "tech job" in body["detail"]
        assert body["domain"]["is_tech"] is False

    def test_tech_jd_is_analysed(self, client, sample_text):
        response = client.post(
            "/api/analyze/job-description",
            json={"resume_text": sample_text, "job_description": ML_JD, "persist": False},
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["job_description_analysis"]["domain"]["is_tech"] is True
        assert "fit_assessment" in payload


# ---------------------------------------------------------------------------
# The fit verdict
# ---------------------------------------------------------------------------


class TestFitVerdict:
    def test_strong_resume_against_a_matching_jd_is_a_fit(self, parsed_strong):
        requirement, analysis = _jd_analysis(ML_JD)
        payload = get_analysis_service().run(parsed_strong, requirement, jd_analysis=analysis)
        fit = payload["fit_assessment"]
        assert fit["verdict"] in {VERDICT_STRONG, VERDICT_GOOD}, fit["headline"]
        assert fit["is_fit"] is True
        assert not fit["knockouts"]

    def test_junior_resume_against_a_six_year_jd_is_not_a_fit(self, parsed):
        requirement, analysis = _jd_analysis(SENIOR_ML_JD)
        assert analysis["years_required"] >= 6
        payload = get_analysis_service().run(parsed, requirement, jd_analysis=analysis)
        fit = payload["fit_assessment"]

        assert fit["verdict"] == VERDICT_NONE
        assert fit["is_fit"] is False
        assert "experience" in fit["knockouts"]
        experience = next(c for c in fit["checks"] if c["key"] == "experience")
        assert experience["knockout"] is True
        assert experience["status"] == "fail"

    def test_catalog_analysis_never_knocks_out_on_experience(self, parsed):
        """Our own level expectation is guidance; only the employer's number filters."""
        payload = get_analysis_service().analyse_target(
            parsed, "google", "software-engineer", "senior"
        )
        experience = next(
            c for c in payload["fit_assessment"]["checks"] if c["key"] == "experience"
        )
        assert experience["knockout"] is False

    def test_every_check_reports_its_own_reason(self, parsed):
        payload = get_analysis_service().analyse_target(
            parsed, "amazon", "machine-learning-engineer", "entry"
        )
        fit = payload["fit_assessment"]
        assert {c["key"] for c in fit["checks"]} == {
            "required_coverage",
            "blocking_gaps",
            "experience",
            "degree",
            "overall_score",
        }
        for check in fit["checks"]:
            assert check["status"] in {"pass", "warn", "fail"}
            assert check["detail"].strip()
        assert 0 <= fit["confidence"] <= 100
        assert 0 <= fit["required_coverage"] <= 100

    def test_a_knockout_always_forces_not_a_fit(self, parsed):
        requirement, analysis = _jd_analysis(SENIOR_ML_JD)
        fit = get_analysis_service().run(parsed, requirement, jd_analysis=analysis)[
            "fit_assessment"
        ]
        assert fit["knockouts"]
        assert fit["verdict"] == VERDICT_NONE

    def test_coverage_below_the_knockout_line_fails(self, parsed):
        """A target the resume barely covers must not be sold as partially fine.

        The sample resume covers ~30% of a senior data-engineering role. The
        coverage assertion is unconditional on purpose: if the datasets drift
        far enough that this target stops being a knockout, the test should
        fail loudly rather than quietly stop checking anything.
        """
        payload = get_analysis_service().analyse_target(
            parsed, "netflix", "data-engineer", "senior"
        )
        fit = payload["fit_assessment"]
        assert fit["required_coverage"] / 100 < COVERAGE_KNOCKOUT
        assert "required_coverage" in fit["knockouts"]
        assert fit["verdict"] == VERDICT_NONE

    def test_middle_band_is_reachable(self, parsed):
        """Coverage can clear the Good bar and still be held back by the score."""
        fit = get_analysis_service().analyse_target(
            parsed, "nvidia", "machine-learning-engineer", "senior"
        )["fit_assessment"]
        assert not fit["knockouts"]
        assert fit["verdict"] not in {VERDICT_STRONG}
        assert 0 < fit["required_coverage"] < 100

    def test_verdict_leads_the_insights(self, parsed):
        payload = get_analysis_service().analyse_target(
            parsed, "amazon", "machine-learning-engineer", "entry"
        )
        assert payload["insights"][0]["type"] == "fit"

    def test_next_step_addresses_the_knockout_not_cosmetics(self, parsed):
        requirement, analysis = _jd_analysis(SENIOR_ML_JD)
        fit = get_analysis_service().run(parsed, requirement, jd_analysis=analysis)[
            "fit_assessment"
        ]
        # With a hard filter in the way, advice about bullet wording is noise.
        assert "bullet" not in fit["next_step"].lower()


# ---------------------------------------------------------------------------
# The builder's consent modes
# ---------------------------------------------------------------------------


class TestBuilderModes:
    def test_confirmed_mode_adds_nothing_unticked(self, parsed, requirement):
        result = get_resume_builder().build(parsed, requirement, BuildOptions())
        assert result["mode"] == MODE_CONFIRMED
        assert result["skills_added"]["confirmed"] == []
        assert result["skills_added"]["auto_added"] == []
        assert result["auto_added_skills"] == []
        assert result["auto_add_warning"] == ""

    def test_confirmed_mode_adds_only_what_was_ticked(self, parsed, requirement):
        builder = get_resume_builder()
        baseline = builder.build(parsed, requirement, BuildOptions())
        missing = [
            m["skill"]
            for m in _missing_skills(parsed, requirement)
            if m["skill"] not in baseline["text"]
        ]
        assert missing, "the sample resume should have gaps to tick"
        ticked = missing[:1]

        result = builder.build(
            parsed, requirement, BuildOptions(confirmed_skills=ticked)
        )
        assert result["skills_added"]["confirmed"] == ticked
        # Nothing beyond the tick arrived.
        assert not result["skills_added"]["auto_added"]
        for skill in missing[1:]:
            assert skill not in result["skills_added"]["confirmed"]

    def test_auto_mode_adds_every_gap_and_labels_them(self, parsed, requirement):
        result = get_resume_builder().build(
            parsed, requirement, BuildOptions(mode=MODE_AUTO_ADD)
        )
        assert result["mode"] == MODE_AUTO_ADD
        added = result["auto_added_skills"]
        assert added, "auto mode must add the missing skills"
        assert result["skills_added"]["auto_added"] == added
        assert "without your confirmation" in result["auto_add_warning"]

        expected = {m["skill"] for m in _missing_skills(parsed, requirement)}
        assert set(added) == expected

        # The change log has to say it happened, not just the payload.
        assert any("Auto-added" in change["action"] for change in result["changes"])

    def test_auto_mode_is_never_the_default(self):
        assert BuildOptions().normalised_mode() == MODE_CONFIRMED
        assert BuildOptions(mode="nonsense").normalised_mode() == MODE_CONFIRMED

    def test_learning_skills_stay_out_of_the_skills_list(self, parsed, requirement):
        missing = [m["skill"] for m in _missing_skills(parsed, requirement)]
        learning = missing[:2]
        result = get_resume_builder().build(
            parsed, requirement, BuildOptions(learning_skills=learning)
        )
        assert result["skills_added"]["learning"] == learning
        listed = {s for group in result["resume"]["skills"] for s in group["skills"]}
        for skill in learning:
            assert skill not in listed
        assert "Currently learning:" in result["text"]

    def test_contact_overrides_and_missing_details(self, parsed, requirement):
        result = get_resume_builder().build(
            parsed, requirement, BuildOptions(details={"phone": "+91 90000 11111"})
        )
        assert "+91 90000 11111" in result["text"]
        assert "phone" not in {d["field"] for d in result["missing_details"]}
        for entry in result["missing_details"]:
            assert entry["label"]
            assert isinstance(entry["essential"], bool)


# ---------------------------------------------------------------------------
# Honesty
# ---------------------------------------------------------------------------


class TestBuilderHonesty:
    @pytest.mark.parametrize("fixture", ["parsed", "parsed_strong"])
    def test_no_invented_numbers_in_confirmed_mode(self, fixture, requirement, request):
        """The generated document must not contain a number the resume never stated.

        Checked against the *shipped* text: bracketed blanks are prompts to the
        user, are stripped on export, and their example wording ("F1") is not a
        claim about the candidate.
        """
        resume = request.getfixturevalue(fixture)
        source = resume.raw_text
        built = get_resume_builder().build(resume, requirement, BuildOptions())
        shipped = strip_placeholders(built["text"])

        assert not _introduces_metric(source, shipped), _new_numbers(source, shipped)

    def test_auto_mode_only_introduces_digits_from_skill_names(self, parsed, requirement):
        """Auto mode may add "Amazon S3" - a proper noun, not a metric."""
        built = get_resume_builder().build(
            parsed, requirement, BuildOptions(mode=MODE_AUTO_ADD)
        )
        shipped = strip_placeholders(built["text"])
        added = " ".join(built["auto_added_skills"])
        for number in _new_numbers(parsed.raw_text, shipped):
            assert number in added, f"{number} came from neither the resume nor a skill name"

    def test_support_verbs_are_not_promoted(self, parsed, requirement):
        """"Helped the data team" must never become "Led the data team"."""
        built = get_resume_builder().build(parsed, requirement, BuildOptions())
        text = built["text"].lower()
        assert "helped the data team" in text
        assert "led the data team" not in text

    def test_metric_placeholders_survive_into_the_text(self, parsed, requirement):
        built = get_resume_builder().build(parsed, requirement, BuildOptions())
        assert built["placeholder_count"] > 0
        assert "[add a measurable result" in built["text"]
        assert built["verification_note"]

    def test_score_never_drops_with_no_claims_added(self, parsed, requirement):
        """Presentation-only changes must not cost the candidate points."""
        result = get_resume_builder().build(parsed, requirement, BuildOptions())
        assert result["score"]["after"] >= result["score"]["before"]

    def test_before_and_after_are_measured_against_the_same_target(
        self, parsed, requirement
    ):
        result = get_resume_builder().build(parsed, requirement, BuildOptions())
        for key in ("before", "after"):
            assert 0 <= result["score"][key] <= 100
        assert result["score"]["fit_before"]
        assert result["score"]["fit_after"]

    def test_rebuilt_requirement_round_trips(self, requirement):
        """A JD target lives only in the stored payload, so this must be exact."""
        restored = JobRequirement.from_dict(requirement.to_dict())
        assert restored.requirement_text() == requirement.requirement_text()
        assert len(restored.by_tier("required")) == len(requirement.by_tier("required"))
        assert restored.weights == requirement.weights


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------


class TestExport:
    @pytest.fixture(scope="class")
    def built_text(self, parsed, requirement):
        return get_resume_builder().build(parsed, requirement, BuildOptions())["text"]

    def test_docx_opens_and_is_ats_plain(self, built_text):
        from docx import Document

        data = render(built_text, "docx")
        assert data[:2] == b"PK"  # a zip container, as DOCX must be

        document = Document(io.BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        assert paragraphs
        # Tables and images are the two things ATS parsers mangle.
        assert len(document.tables) == 0
        assert any(p.isupper() for p in paragraphs), "section headings should be present"

    def test_pdf_is_a_real_pdf(self, built_text):
        data = render(built_text, "pdf")
        assert data.startswith(b"%PDF")

        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as document:
            assert document.page_count >= 1
            assert "PROFESSIONAL SUMMARY" in document[0].get_text()

    def test_txt_round_trips(self, built_text):
        data = render(built_text, "txt")
        assert data.decode("utf-8").strip()

    def test_placeholders_are_stripped_by_default(self, built_text):
        assert "[add a measurable result" in built_text
        assert "[add a measurable result" not in render(built_text, "txt").decode("utf-8")
        kept = render(built_text, "txt", strip=False).decode("utf-8")
        assert "[add a measurable result" in kept

    def test_stripping_leaves_no_dangling_punctuation(self, built_text):
        cleaned = strip_placeholders(built_text)
        for line in cleaned.split("\n"):
            assert not line.rstrip().endswith(","), line
            assert line.strip() != "-", "an empty bullet was left behind"

    def test_unsupported_format_is_rejected(self, built_text):
        from services.resume_export import ResumeExportError

        with pytest.raises(ResumeExportError):
            render(built_text, "rtf")


class TestBuilderAPI:
    @pytest.fixture(scope="class")
    def analysis_id(self, client, sample_text):
        upload = client.post("/api/resume/parse-text", data={"resume_text": sample_text})
        resume_id = upload.json()["resume_id"]
        analysis = client.post(
            "/api/analyze",
            json={
                "resume_id": resume_id,
                "company": "amazon",
                "role": "machine-learning-engineer",
                "level": "entry",
            },
        )
        return analysis.json()["analysis_id"]

    def test_build_endpoint(self, client, analysis_id):
        response = client.post("/api/resume/build", json={"analysis_id": analysis_id})
        assert response.status_code == 200
        body = response.json()
        assert body["mode"] == MODE_CONFIRMED
        assert body["text"].strip()
        assert body["score"]["after"] >= body["score"]["before"]

    def test_build_rejects_a_skill_that_is_not_a_gap(self, client, analysis_id):
        response = client.post(
            "/api/resume/build",
            json={"analysis_id": analysis_id, "confirmed_skills": ["Python"]},
        )
        assert response.status_code == 422
        assert "not gaps in this analysis" in response.json()["detail"]

    def test_build_rejects_confirmed_and_learning_overlap(self, client, analysis_id):
        gaps = client.get(f"/api/analysis/{analysis_id}").json()["skill_gap"]
        skill = gaps["missing_high_priority"][0]["skill"]
        response = client.post(
            "/api/resume/build",
            json={
                "analysis_id": analysis_id,
                "confirmed_skills": [skill],
                "learning_skills": [skill],
            },
        )
        assert response.status_code == 422
        assert "cannot be both" in response.json()["detail"]

    def test_build_404s_on_an_unknown_analysis(self, client):
        assert client.post("/api/resume/build", json={"analysis_id": 999999}).status_code == 404

    @pytest.mark.parametrize(
        "fmt,prefix,content_type",
        [
            ("docx", b"PK", "officedocument"),
            ("pdf", b"%PDF", "application/pdf"),
            ("txt", b"", "text/plain"),
        ],
    )
    def test_export_endpoint(self, client, analysis_id, fmt, prefix, content_type):
        built = client.post("/api/resume/build", json={"analysis_id": analysis_id}).json()
        response = client.post(
            "/api/resume/export",
            json={"text": built["text"], "format": fmt, "file_name": "Test Candidate"},
        )
        assert response.status_code == 200
        assert response.content.startswith(prefix)
        assert content_type in response.headers["content-type"]
        assert "attachment" in response.headers["content-disposition"]
        assert f".{fmt}" in response.headers["content-disposition"]

    def test_export_rejects_empty_text(self, client):
        assert client.post(
            "/api/resume/export", json={"text": "too short", "format": "pdf"}
        ).status_code == 422


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _missing_skills(resume, requirement) -> list[dict]:
    """The gap list as the builder sees it: missing required + preferred."""
    from services.job_matcher import get_job_matcher
    from services.skill_extractor import get_skill_extractor

    skills = get_skill_extractor().extract(resume)
    match = get_job_matcher().match(resume, skills, requirement)
    return [
        m.to_dict()
        for m in match.matches
        if m.status == "missing" and m.tier in {"required", "preferred"}
    ]


def _new_numbers(original: str, generated: str) -> list[str]:
    originals = set(_NUMBER_RE.findall(original))
    return sorted({n for n in _NUMBER_RE.findall(generated) if n not in originals})
