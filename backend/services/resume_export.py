"""Resume export.

Turns the builder's plain text - after the user has edited it - into a DOCX,
PDF or TXT download.

Both binary formats are deliberately ATS-plain: one column, black text, no
tables, no images, no colour, no text boxes. Applicant tracking systems parse
those structures badly or not at all, and a resume that reads beautifully but
imports as gibberish is worse than a dull one that imports cleanly.

The text is interpreted with the same conventions the builder writes:

    ANANYA SHARMA            first line          -> name
    a@b.com | +91 …          contact line        -> contact
    TECHNICAL SKILLS         an ALL-CAPS line    -> section heading
    - Built a thing          a "- " line         -> bullet
    Languages: Python        anything else       -> paragraph
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, Iterable

FORMAT_DOCX = "docx"
FORMAT_PDF = "pdf"
FORMAT_TXT = "txt"
FORMATS = (FORMAT_DOCX, FORMAT_PDF, FORMAT_TXT)

CONTENT_TYPES = {
    FORMAT_DOCX: "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    FORMAT_PDF: "application/pdf",
    FORMAT_TXT: "text/plain; charset=utf-8",
}

# Two kinds of blank, with very different consequences on export.
#
# "[add a measurable result …]" is an optional prompt. Deleting it leaves a
# sentence that is still true - the bullet simply has no number.
#
# "<model type>", "<latency>" are *required* blanks: the surrounding sentence
# only makes sense once they are filled. Deleting those would turn
# "Deployed a <model type> model … at <latency> ms p95" into
# "Deployed a model … at ms p95" - a broken sentence that also claims work the
# candidate may never have done. So they block the export instead.
_PROMPT_RE = re.compile(r"\s*\[[^\]]*\]")
_BLANK_RE = re.compile(r"<[^>]*>")

# The one <…> blank that is safe to remove: the sentence reads fine without it,
# provided the dangling connector goes too.
_TOOL_BLANK = "<the specific tools you used>"
_TOOL_BLANK_RE = re.compile(
    r"\s*(?:\b(?:using|with|via|through)\b\s*)?" + re.escape(_TOOL_BLANK), re.IGNORECASE
)

_TRAILING_PUNCT_RE = re.compile(r"[ ,;:]+(?=\.?$)")

_HEADING_MAX_WORDS = 6


class ResumeExportError(ValueError):
    """Raised when a document cannot be produced."""


@dataclass
class Line:
    kind: str  # name | contact | heading | bullet | text
    text: str


def required_blanks(text: str) -> list[str]:
    """``<...>`` blanks the user still has to supply before this can be sent."""
    return [
        blank
        for blank in _BLANK_RE.findall(text or "")
        if blank.lower() != _TOOL_BLANK.lower()
    ]


def strip_placeholders(text: str) -> str:
    """Remove the optional prompts and tidy what is left behind.

    Required ``<...>`` blanks are deliberately left in place - see the note
    above ``_PROMPT_RE``. Anything about to produce a document should check
    :func:`required_blanks` first.
    """
    out: list[str] = []
    for line in text.split(chr(10)):
        cleaned = _TOOL_BLANK_RE.sub("", line)
        cleaned = _PROMPT_RE.sub("", cleaned)
        cleaned = _TRAILING_PUNCT_RE.sub("", cleaned).rstrip()
        # A bullet that was *only* a placeholder is dropped, not left empty.
        if cleaned.strip() in {"-", "-.", ""} and line.strip():
            continue
        out.append(cleaned)
    return chr(10).join(out)

def parse_lines(text: str) -> list[Line]:
    """Classify each line so every renderer formats from the same structure."""
    lines: list[Line] = []
    seen_content = False
    for raw in text.split("\n"):
        stripped = raw.strip()
        if not stripped:
            lines.append(Line("blank", ""))
            continue

        if not seen_content:
            seen_content = True
            lines.append(Line("name", stripped))
            continue

        if not any(l.kind == "contact" for l in lines) and _looks_like_contact(stripped, lines):
            lines.append(Line("contact", stripped))
            continue

        if _is_heading(stripped):
            lines.append(Line("heading", stripped))
        elif stripped.startswith(("- ", "• ", "* ")):
            lines.append(Line("bullet", stripped[2:].strip()))
        else:
            lines.append(Line("text", stripped))
    return lines


def _looks_like_contact(line: str, seen: list[Line]) -> bool:
    """The contact line is the one right under the name, with contact markers."""
    if not seen or seen[-1].kind != "name":
        return False
    return bool(re.search(r"@|\||https?://|\+\d|linkedin|github", line, re.IGNORECASE))


def _is_heading(line: str) -> bool:
    letters = [c for c in line if c.isalpha()]
    if not letters or len(line.split()) > _HEADING_MAX_WORDS:
        return False
    return all(c.isupper() for c in letters)


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------


def render(text: str, fmt: str, *, strip: bool = True) -> bytes:
    if fmt not in FORMATS:
        raise ResumeExportError(f"Unsupported format '{fmt}'. Use one of: {', '.join(FORMATS)}.")
    if strip:
        outstanding = required_blanks(text)
        if outstanding:
            shown = ", ".join(dict.fromkeys(outstanding))[:160]
            raise ResumeExportError(
                f"{len(outstanding)} blank(s) still need your own details: {shown}. "
                "Fill them in, or delete the line - removing them for you would put a "
                "claim on your resume that you have not backed up."
            )

    body = strip_placeholders(text) if strip else text
    if not body.strip():
        raise ResumeExportError("There is no resume text to export.")

    if fmt == FORMAT_TXT:
        return body.encode("utf-8")
    if fmt == FORMAT_DOCX:
        return _render_docx(parse_lines(body))
    return _render_pdf(parse_lines(body))


def _render_docx(lines: Iterable[Line]) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - dependency is in requirements
        raise ResumeExportError(
            "DOCX export needs python-docx. Install it with: pip install python-docx"
        ) from exc

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0, 0, 0)

    for section in document.sections:
        section.top_margin = section.bottom_margin = Pt(40)
        section.left_margin = section.right_margin = Pt(46)

    def paragraph(text: str, *, size: float, bold: bool = False, space_after: float = 3,
                  align=None, style_name: str | None = None):
        para = document.add_paragraph(style=style_name)
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(0)
        if align is not None:
            para.alignment = align
        run = para.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return para

    for line in lines:
        if line.kind == "blank":
            continue
        if line.kind == "name":
            paragraph(line.text, size=19, bold=True, space_after=2,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line.kind == "contact":
            paragraph(line.text, size=9.5, space_after=10,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line.kind == "heading":
            para = paragraph(line.text, size=11.5, bold=True, space_after=4)
            para.paragraph_format.space_before = Pt(10)
            _add_bottom_border(para, OxmlElement, qn)
        elif line.kind == "bullet":
            paragraph(line.text, size=10.5, space_after=2, style_name="List Bullet")
        else:
            paragraph(line.text, size=10.5, space_after=2)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_bottom_border(paragraph: Any, OxmlElement: Any, qn: Any) -> None:
    """A single black rule under a section heading.

    Done through the XML because python-docx exposes no paragraph-border API.
    A border is safe for ATS parsing in a way a table or text box is not.
    """
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


_PDF_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; color: #000000; line-height: 1.38; }
h1 { font-size: 19pt; margin: 0 0 2pt 0; text-align: center; }
p.contact { font-size: 9.5pt; margin: 0 0 12pt 0; text-align: center; }
h2 { font-size: 11.5pt; margin: 11pt 0 3pt 0; border-bottom: 1px solid #000000; }
p { margin: 0 0 3pt 0; }
ul { margin: 0 0 3pt 0; padding-left: 13pt; }
li { margin: 0 0 2pt 0; }
"""


def _render_pdf(lines: list[Line]) -> bytes:
    try:
        import pymupdf as fitz
    except ImportError:  # pragma: no cover - fallback for older installs
        try:
            import fitz  # type: ignore
        except ImportError as exc:
            raise ResumeExportError(
                "PDF export needs PyMuPDF. Install it with: pip install pymupdf"
            ) from exc

    html = ["<html><body>"]
    in_list = False
    for line in lines:
        if line.kind == "bullet":
            if not in_list:
                html.append("<ul>")
                in_list = True
            html.append(f"<li>{_escape(line.text)}</li>")
            continue
        if in_list:
            html.append("</ul>")
            in_list = False
        if line.kind == "name":
            html.append(f"<h1>{_escape(line.text)}</h1>")
        elif line.kind == "contact":
            html.append(f'<p class="contact">{_escape(line.text)}</p>')
        elif line.kind == "heading":
            html.append(f"<h2>{_escape(line.text)}</h2>")
        elif line.kind == "text":
            html.append(f"<p>{_escape(line.text)}</p>")
    if in_list:
        html.append("</ul>")
    html.append("</body></html>")

    story = fitz.Story(html="".join(html), user_css=_PDF_CSS)
    buffer = io.BytesIO()
    writer = fitz.DocumentWriter(buffer)
    page = fitz.paper_rect("a4")
    area = page + (52, 48, -52, -48)

    more = True
    guard = 0
    while more and guard < 20:  # guard against a pathological layout loop
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
        guard += 1
    writer.close()
    return buffer.getvalue()


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def safe_file_name(name: str, fmt: str) -> str:
    """A download name that cannot escape the filename slot."""
    stem = re.sub(r"[^A-Za-z0-9 _.-]", "", (name or "").strip()) or "resume"
    stem = re.sub(r"\s+", "_", stem)[:80]
    if stem.lower().endswith(f".{fmt}"):
        return stem
    return f"{stem}.{fmt}"
